from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

# Create a presentation object
prs = Presentation()

# Slide 1: Title Slide
slide = prs.slides.add_slide(prs.slide_layouts[0])
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "The Joy of Traveling"
subtitle.text = "Explore, Experience, Enjoy!"

# Slide 2: Why Do People Travel?
slide = prs.slides.add_slide(prs.slide_layouts[1])
title = slide.shapes.title
title.text = "Why Do People Travel?"
content = slide.placeholders[1]
content.text = "• To explore new places\n• To experience different cultures\n• For relaxation and leisure\n• To learn and grow"

# Slide 3: Popular Travel Destinations
slide = prs.slides.add_slide(prs.slide_layouts[1])
title = slide.shapes.title
title.text = "Popular Travel Destinations"
content = slide.placeholders[1]
content.text = ("• Europe: Paris, Rome, London\n"
                "• Asia: Tokyo, Bangkok, Bali\n"
                "• North America: New York, Vancouver, Cancun\n"
                "• Natural Wonders: Grand Canyon, Great Barrier Reef, Amazon Rainforest")

# Slide 4: Types of Travel Experiences
slide = prs.slides.add_slide(prs.slide_layouts[1])
title = slide.shapes.title
title.text = "Types of Travel Experiences"
content = slide.placeholders[1]
content.text = ("• Adventure Travel (e.g., hiking, diving)\n"
                "• Cultural Travel (e.g., museums, festivals)\n"
                "• Luxury Travel (e.g., resorts, fine dining)\n"
                "• Eco-tourism (e.g., wildlife conservation, nature hikes)")

# Slide 5: Benefits of Traveling
slide = prs.slides.add_slide(prs.slide_layouts[1])
title = slide.shapes.title
title.text = "Benefits of Traveling"
content = slide.placeholders[1]
content.text = ("• Broadens perspective and understanding\n"
                "• Improves mental well-being and reduces stress\n"
                "• Boosts creativity and inspiration\n"
                "• Builds confidence and adaptability")

# Slide 6: Planning Your Trip
slide = prs.slides.add_slide(prs.slide_layouts[1])
title = slide.shapes.title
title.text = "Planning Your Trip"
content = slide.placeholders[1]
content.text = ("• Decide on a destination\n"
                "• Research accommodation and transportation\n"
                "• Create a budget and save up\n"
                "• Plan activities and pack accordingly")

# Slide 7: Traveling on a Budget
slide = prs.slides.add_slide(prs.slide_layouts[1])
title = slide.shapes.title
title.text = "Traveling on a Budget"
content = slide.placeholders[1]
content.text = ("• Look for affordable destinations\n"
                "• Use budget airlines and off-peak flights\n"
                "• Stay in hostels or Airbnbs\n"
                "• Try local street food and markets")

# Slide 8: Travel Safety Tips
slide = prs.slides.add_slide(prs.slide_layouts[1])
title = slide.shapes.title
title.text = "Travel Safety Tips"
content = slide.placeholders[1]
content.text = ("• Keep your valuables secure\n"
                "• Stay aware of local customs and laws\n"
                "• Use reputable transportation options\n"
                "• Have copies of important documents")

# Slide 9: Travel Essentials Checklist
slide = prs.slides.add_slide(prs.slide_layouts[1])
title = slide.shapes.title
title.text = "Travel Essentials Checklist"
content = slide.placeholders[1]
content.text = ("• Passport and ID\n"
                "• Travel insurance\n"
                "• Comfortable clothing and shoes\n"
                "• Mobile charger and adapter\n"
                "• Basic first-aid kit")

# Slide 10: Memories and Souvenirs
slide = prs.slides.add_slide(prs.slide_layouts[1])
title = slide.shapes.title
title.text = "Memories and Souvenirs"
content = slide.placeholders[1]
content.text = "• Capture the moments\n• Collect souvenirs\n• Keep a travel journal\n• Share experiences with others"

# Save the presentation
prs.save("Travel_Presentation.pptx")

print("Presentation created successfully!")
